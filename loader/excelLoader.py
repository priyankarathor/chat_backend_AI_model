from langchain_core.documents import Document


def _row_to_text(headers, values):
    return "\n".join(
        f"{header}: {value if value is not None else ''}"
        for header, value in zip(headers, values)
    )


def load_excel(file_path):
    if file_path.lower().endswith(".xlsx"):
        return _load_xlsx(file_path)

    return _load_xls(file_path)


def _load_xlsx(file_path):
    from openpyxl import load_workbook

    workbook = load_workbook(file_path, read_only=True, data_only=True)
    documents = []

    try:
        for worksheet in workbook.worksheets:
            rows = worksheet.iter_rows(values_only=True)
            headers = [str(value or f"column_{index}") for index, value in enumerate(next(rows, ()), start=1)]

            for row_number, row in enumerate(rows, start=2):
                documents.append(
                    Document(
                        page_content=_row_to_text(headers, row),
                        metadata={"source": file_path, "sheet": worksheet.title, "row": row_number},
                    )
                )
    finally:
        workbook.close()

    return documents


def _load_xls(file_path):
    import xlrd

    workbook = xlrd.open_workbook(file_path, on_demand=True)
    documents = []

    try:
        for worksheet in workbook.sheets():
            if worksheet.nrows == 0:
                continue

            headers = [str(value or f"column_{index}") for index, value in enumerate(worksheet.row_values(0), start=1)]

            for row_index in range(1, worksheet.nrows):
                documents.append(
                    Document(
                        page_content=_row_to_text(headers, worksheet.row_values(row_index)),
                        metadata={"source": file_path, "sheet": worksheet.name, "row": row_index + 1},
                    )
                )
    finally:
        workbook.release_resources()

    return documents
